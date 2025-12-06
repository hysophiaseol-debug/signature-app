from flask import Flask, request
import base64
from openpyxl import Workbook
from openpyxl.drawing.image import Image
import io

app = Flask(__name__)

@app.route('/upload', methods=['POST'])
def upload():
    data = request.get_json()

    header, encoded = data['signature'].split(',')
    img_binary = base64.b64decode(encoded)

    wb = Workbook()
    ws = wb.active

    ws['A1'] = '제출 시간'
    ws['B1'] = '이름'
    ws['C1'] = '소속'

    ws['A2'] = data['time']
    ws['B2'] = data['name']
    ws['C2'] = data['org']

    img_stream = io.BytesIO(img_binary)
    img = Image(img_stream)
    img.width = 250
    img.height = 120
    ws.add_image(img, 'E2')

    wb.save("submitted.xlsx")

    return {"status": "success"}

@app.route('/', methods=['GET'])
def home():
    return "Signature server is running."

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=10000)
